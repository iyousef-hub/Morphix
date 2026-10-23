"""Real document and image OCR regression checks (no GUI dependency)."""
import sys
from pathlib import Path

import pytest
import pymupdf
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from ocr_engine import extract_pdf_text, searchable_pdf, find_tessdata
from office_engine import find_office_engine, office_to_pdf
from core_ops import archive_files, rar_writer_available
from test_engine import engine_exports, Progress


def test_ocr_models_are_bundled():
    d = find_tessdata()
    assert d is not None
    assert (d/'ara.traineddata').stat().st_size > 20000
    assert (d/'eng.traineddata').stat().st_size > 20000


def test_ocr_real_arabic_scanned_pdf(tmp_path):
    fonts = list(Path('/usr/share/fonts').rglob('NotoSansArabic-Regular.ttf'))
    if not fonts:
        pytest.skip('Linux Arabic test font not installed in this environment')
    image = Image.new('RGB', (1200, 340), 'white')
    draw = ImageDraw.Draw(image)
    arabic = ImageFont.truetype(str(fonts[0]), 66)
    draw.text((120, 85), 'مرحبا', font=arabic, fill='black', direction='rtl')
    image_file = tmp_path/'scan.png'
    image.save(image_file)
    pdf = tmp_path/'scan.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page(width=600, height=170)
        page.insert_image(page.rect, filename=str(image_file))
        doc.save(pdf)
    out = extract_pdf_text([pdf], tmp_path)[0]
    assert 'مرحبا' in out.read_text(encoding='utf-8-sig')
    assert 'Page 1' in out.read_text(encoding='utf-8-sig')


def test_ocr_english_scanned_pdf(tmp_path):
    font_locations = list(Path('/usr/share/fonts').rglob('DejaVuSans.ttf'))
    if not font_locations: pytest.skip('English test font missing')
    image = Image.new('RGB', (1100, 280), 'white')
    ImageDraw.Draw(image).text((80, 70), 'Morphix Lecture 2026',
                               font=ImageFont.truetype(str(font_locations[0]), 59), fill='black')
    png = tmp_path/'en.png'; image.save(png)
    pdf = tmp_path/'en.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page(width=550, height=140)
        page.insert_image(page.rect, filename=str(png))
        doc.save(pdf)
    output = extract_pdf_text([pdf],tmp_path)[0]
    assert 'Morphix Lecture 2026' in output.read_text(encoding='utf-8-sig')


def test_searchable_scanned_pdf(tmp_path):
    fonts = list(Path('/usr/share/fonts').rglob('NotoSansArabic-Regular.ttf'))
    if not fonts:
        pytest.skip('Arabic test font not available')
    image = Image.new('RGB', (1100, 260), 'white')
    ImageDraw.Draw(image).text((80, 60), 'مرحبا', font=ImageFont.truetype(str(fonts[0]), 60),
                               fill='black', direction='rtl')
    png = tmp_path/'img.png'; image.save(png)
    pdf = tmp_path/'orig.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page(width=550, height=130)
        page.insert_image(page.rect, filename=str(png))
        doc.save(pdf)
    out = searchable_pdf([pdf],tmp_path)[0]
    with pymupdf.open(out) as doc:
        assert len(doc) == 1
        assert doc[0].rect.width == 550
        assert 'مرحبا' in doc[0].get_text('text')


def test_native_text_skips_ocr(tmp_path, monkeypatch):
    import ocr_engine
    pdf = tmp_path/'native.pdf'
    with pymupdf.open() as doc:
        doc.new_page().insert_text((40,40), 'Digitally typed sample document')
        doc.save(pdf)
    def forbidden(*args, **kwargs):
        raise AssertionError('OCR must not run on a searchable text page')
    monkeypatch.setattr(ocr_engine, '_ocr_page', forbidden)
    out = extract_pdf_text([pdf], tmp_path)[0]
    assert 'Digitally typed' in out.read_text(encoding='utf-8-sig')


def test_direct_engine_ocr_dispatch(tmp_path):
    pdf = tmp_path/'native.pdf'
    with pymupdf.open() as doc:
        doc.new_page().insert_text((40,40), 'English text in PDF')
        doc.save(pdf)
    outputs = engine_exports()['execute']('extract_pdf_text',[str(pdf)],str(tmp_path),Progress())
    assert outputs[0].suffix=='.txt'
    assert 'English text' in outputs[0].read_text(encoding='utf-8-sig')


@pytest.mark.skipif(find_office_engine() is None, reason='Office engine not installed on CI host')
def test_docx_to_pdf(tmp_path):
    from docx import Document
    docx = tmp_path/'lecture.docx'
    doc = Document(); doc.add_heading('Morphix Office test', 1); doc.add_paragraph('This is an exported Word document.'); doc.save(docx)
    outputs = office_to_pdf([docx], tmp_path)
    with pymupdf.open(outputs[0]) as pdf:
        assert len(pdf)>=1
        assert 'This is an exported Word document' in pdf[0].get_text()


@pytest.mark.skipif(find_office_engine() is None, reason='Office engine not installed on CI host')
def test_pptx_to_pdf(tmp_path):
    from pptx import Presentation
    pptx = tmp_path/'slides.pptx'
    p=Presentation(); slide=p.slides.add_slide(p.slide_layouts[0]); slide.shapes.title.text='Morphix Presentation'; p.save(pptx)
    outputs=office_to_pdf([pptx],tmp_path)
    with pymupdf.open(outputs[0]) as pdf:
        assert len(pdf)==1
        assert 'Morphix Presentation' in pdf[0].get_text()


@pytest.mark.skipif(find_office_engine() is None, reason='Office engine not installed on CI host')
def test_xlsx_to_pdf(tmp_path):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws['A1']='Morphix Table'; ws['A2']=11; wb.save(tmp_path/'data.xlsx')
    outputs=office_to_pdf([tmp_path/'data.xlsx'],tmp_path)
    with pymupdf.open(outputs[0]) as pdf:
        assert len(pdf)>=1
        assert 'Morphix Table' in pdf[0].get_text()


def test_no_fake_rar_without_writer(tmp_path,monkeypatch):
    import core_ops
    monkeypatch.setattr(core_ops,'find_rar_writer',lambda:None)
    p=tmp_path/'a.txt';p.write_text('hello')
    with pytest.raises(RuntimeError,match='licensed RAR writer'):
        archive_files([p],tmp_path,'rar')
    assert not list(tmp_path.glob('*.rar'))


def test_missing_office_writer_reports_clear_error(tmp_path,monkeypatch):
    import office_engine
    monkeypatch.setattr(office_engine,'find_office_engine',lambda:None)
    sample=tmp_path/'a.docx'; sample.write_bytes(b'test')
    with pytest.raises(RuntimeError, match='needs LibreOffice'):
        office_to_pdf([sample],tmp_path)


def test_office_rejects_wrong_extension(tmp_path):
    if find_office_engine() is None: pytest.skip('Office engine not installed')
    source=tmp_path/'hello.txt';source.write_text('hi')
    with pytest.raises(ValueError,match='Unsupported Office type'):
        office_to_pdf([source],tmp_path)
