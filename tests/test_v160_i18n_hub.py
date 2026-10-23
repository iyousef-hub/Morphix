"""v1.6: bilingual interface, Action Hub shapes/positions, PDF compression levels."""
import os
import re
import runpy
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _strings():
    return runpy.run_path(str(ROOT / 'tools' / 'extract_strings.py'))['STRINGS']


def test_every_interface_string_has_arabic():
    from i18n_ar import AR
    missing = sorted(s for s in _strings() if s not in AR)
    assert not missing, 'Missing Arabic for:\n' + '\n'.join(missing)


def test_placeholders_match():
    from i18n_ar import AR
    ph = re.compile(r'\{(\w+)\}')
    bad = [k for k, v in AR.items() if set(ph.findall(k)) != set(ph.findall(v))]
    assert not bad, bad


def test_tr_switches_language(monkeypatch):
    import i18n
    i18n.set_language('ar', save=False)
    assert i18n.tr('Cancel') == 'إلغاء'
    assert i18n.tr('{n} files', n=3) == '3 ملف'
    assert i18n.is_rtl() and i18n.start_side() == 'right'
    i18n.set_language('en', save=False)
    assert i18n.tr('Cancel') == 'Cancel' and i18n.start_side() == 'left'
    assert i18n.tr('Some unknown text') == 'Some unknown text'


def test_place_rect_positions():
    from PySide6.QtCore import QRect, QPoint
    from hub_panel import place_rect
    avail = QRect(0, 25, 1440, 875)
    w, h = 300, 100
    assert place_rect('top', w, h, QPoint(5, 5), avail) == (569, 39)
    x, y = place_rect('bottom', w, h, QPoint(5, 5), avail)
    assert y + h == avail.bottom() - 14 + 1 or y + h <= avail.bottom()
    assert place_rect('left', w, h, QPoint(0, 0), avail)[0] == 14
    assert place_rect('right', w, h, QPoint(0, 0), avail)[0] + w <= avail.right() + 1
    # cursor near a corner is clamped on-screen
    assert place_rect('cursor', w, h, QPoint(1439, 899), avail) == (1140, 800)
    assert place_rect('remember', w, h, QPoint(0, 0), avail, remembered=(700, 400)) == (550, 350)


def test_pdf_compression_monotonic_and_extreme(tmp_path):
    import io
    import pymupdf as fitz
    from PIL import Image
    import numpy as np
    from compression_engine import save_pdf_compressed, raster_params
    doc = fitz.open()
    rng = np.random.default_rng(1)
    for i in range(3):
        arr = (rng.random((1400, 1000, 3)) * 50 + 170).astype('uint8')
        b = io.BytesIO(); Image.fromarray(arr).save(b, 'JPEG', quality=93)
        pg = doc.new_page(width=612, height=792)
        pg.insert_image(pg.rect, stream=b.getvalue())
    src = tmp_path / 'scan.pdf'
    doc.save(src)
    sizes = {}
    for strength in (10, 45, 80, 100):
        with fitz.open(src) as d:
            out = save_pdf_compressed(d, tmp_path / f'o{strength}.pdf', strength=strength)
        sizes[strength] = out.stat().st_size
        with fitz.open(out) as check:
            assert check.page_count == 3
    assert sizes[100] < sizes[80] <= sizes[45] <= sizes[10] <= src.stat().st_size
    assert sizes[100] < src.stat().st_size / 20
    assert raster_params(85) is None and raster_params(100)['dpi'] < raster_params(86)['dpi']


def test_text_pdf_never_grows(tmp_path):
    import pymupdf as fitz
    from compression_engine import save_pdf_compressed
    doc = fitz.open()
    for i in range(4):
        doc.new_page().insert_text((72, 72), 'Plain text page %d ' % i * 20)
    src = tmp_path / 'text.pdf'
    doc.save(src)
    with fitz.open(src) as d:
        out = save_pdf_compressed(d, tmp_path / 'x.pdf', strength=100)
    assert out.stat().st_size <= src.stat().st_size


def test_ocr_offered_in_pdf_hub_defaults():
    import ast
    tree = ast.parse((ROOT / 'quickwheel.py').read_text(encoding='utf-8'))
    profiles = next(ast.literal_eval(n.value) for n in tree.body
                    if isinstance(n, ast.Assign) and getattr(n.targets[0], 'id', '') == 'PROFILES')
    assert profiles['pdf_single'][1][0] == 'ocr_searchable_pdf'
    assert 'ocr_searchable_pdf' in profiles['pdf_multi'][1]


def test_recorder_swaps_cmd_ctrl_on_mac(monkeypatch):
    import shortcut_recorder as sr
    from PySide6.QtCore import Qt
    monkeypatch.setattr(sr, 'MAC', True)
    assert sr._qt_modifiers(Qt.ControlModifier | Qt.ShiftModifier) == {'<cmd>', '<shift>'}
    assert sr._qt_modifiers(Qt.MetaModifier) == {'<ctrl>'}
    monkeypatch.setattr(sr, 'MAC', False)
    assert sr._qt_modifiers(Qt.ControlModifier) == {'<ctrl>'}


def test_mixed_documents_profile_and_actions(tmp_path):
    from file_intelligence import selection_profile
    from main import supported_actions
    files = []
    for name in ('a.docx', 'b.pptx', 'c.xlsx', 'd.png'):
        p = tmp_path / name
        p.write_bytes(b'x' * 10)
        files.append(str(p))
    assert selection_profile(files) == 'docs_mix'
    assert supported_actions(files)[:2] == ['docs_merge_pdf', 'docs_pdf_each']
    assert selection_profile(files[2:3]) == 'excel_single'
    acts = supported_actions(files[2:3])
    assert {'office_pdf', 'xlsx_pdf_fit', 'xlsx_csv', 'xlsx_split'} <= set(acts)


def test_excel_tools_keep_data(tmp_path):
    import openpyxl
    from openpyxl.styles import Font
    import docs_engine as de
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'المبيعات'
    ws.append(['الاسم', 'الهاتف'])
    ws['A1'].font = Font(bold=True)
    ws.append(['أحمد', '0501234567'])
    wb.create_sheet('Two').append(['x', 1])
    src = tmp_path / 'book.xlsx'
    wb.save(src)
    out = tmp_path / 'out'
    out.mkdir()
    csvs = de.excel_to_csv([str(src)], out)
    assert len(csvs) == 2 and 'أحمد' in csvs[0].read_text(encoding='utf-8-sig')
    parts = de.excel_split_sheets([str(src)], out)
    assert [openpyxl.load_workbook(p).sheetnames for p in parts] == [['المبيعات'], ['Two']]
    merged = de.excel_merge([str(src), str(src)], out)[0]
    m = openpyxl.load_workbook(merged)
    assert len(m.sheetnames) == 4 and m.worksheets[0]['A1'].font.b
    c = tmp_path / 'ids.csv'
    c.write_text('name;phone\nأحمد;0501234567\n', encoding='utf-8')
    x = openpyxl.load_workbook(de.csv_to_excel([str(c)], out)[0]).active
    assert x['B2'].value == '0501234567' and x.sheet_view.rightToLeft


def test_background_keeps_pixels(tmp_path, monkeypatch):
    import numpy as np
    from PIL import Image
    import background_engine as be

    class FakeSession:
        def __init__(self, q):
            pass

        def mask(self, rgb):
            m = Image.new('L', rgb.size, 0)
            m.paste(255, (10, 10, 50, 50))
            return m
    monkeypatch.setattr(be, '_Session', FakeSession)
    monkeypatch.setattr(be, 'model_path', lambda q: tmp_path)
    src = tmp_path / 'p.jpg'
    arr = (np.random.default_rng(0).random((80, 120, 3)) * 255).astype('uint8')
    Image.fromarray(arr).save(src, quality=95)
    out = be.remove_background([str(src)], tmp_path / 'o')[0]
    res = Image.open(out)
    assert res.mode == 'RGBA' and res.size == (120, 80)
    assert (np.asarray(res)[..., :3] == np.asarray(Image.open(src).convert('RGB'))).all()
    assert np.asarray(res)[20, 20, 3] == 255 and np.asarray(res)[5, 5, 3] == 0


def test_windows_words_replace_mac_terms():
    import i18n
    assert i18n.for_windows('Select one or more files in Finder, then try again.') == \
        'Select one or more files in File Explorer, then try again.'
    assert 'Win + Shift' in i18n.for_windows('Examples: ⌘ Command + Shift + M')
    assert 'مستكشف الملفات' in i18n.for_windows('حدّد ملفًا في Finder', 'ar')
    assert i18n.for_windows('Run Launch_Morphix_Mac.command again') == 'Run Morphix setup again'


def test_win_compat_is_noop_off_windows():
    import subprocess, win_compat
    win_compat.apply()
    assert not getattr(subprocess.Popen, '_morphix_hidden', False)
