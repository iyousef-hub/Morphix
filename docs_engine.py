"""Documents: Word / PowerPoint / Excel / images / PDF → PDF (merged or one by one),
plus practical Excel tools. Everything runs locally; originals are never changed.
"""
from __future__ import annotations

import csv
import shutil
import tempfile
from pathlib import Path

from core_ops import unique_path

IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.gif'}
SHEET_EXTS = {'.xlsx', '.xlsm', '.xls', '.ods', '.csv'}
_LANG = 'en'


def _t(text, **fmt):
    try:
        from i18n import trl
        return trl(_LANG, text, **fmt)
    except Exception:
        return text.format(**fmt) if fmt else text


def _emit(progress, value):
    if progress:
        progress(int(value))


# --------------------------------------------------------------------------- PDF
def _to_pdf_parts(paths, workdir, progress=None, span=(0, 80)):
    """Return one PDF (or image) path per input, in the same order."""
    from office_engine import office_to_pdf, OFFICE_EXTS
    parts = []
    for i, p in enumerate(paths, 1):
        src = Path(p)
        ext = src.suffix.lower()
        if ext == '.pdf' or ext in IMAGE_EXTS:
            parts.append(src)
        elif ext in OFFICE_EXTS or ext in SHEET_EXTS:
            parts.append(office_to_pdf([src], workdir)[0])
        else:
            raise ValueError(_t('Unsupported file: {name}', name=src.name))
        _emit(progress, span[0] + (span[1] - span[0]) * i / len(paths))
    return parts


def docs_merge_pdf(paths, out_dir, progress=None, ui_lang='en'):
    """Everything (Word, PowerPoint, Excel, images, PDF) → one PDF, in the given order."""
    global _LANG
    _LANG = ui_lang
    from pdf_merge import merge_documents

    class _P:
        def emit(self, v):
            pass
    with tempfile.TemporaryDirectory(prefix='morphix_docs_') as tmp:
        parts = _to_pdf_parts(paths, tmp, progress)
        doc = merge_documents([str(p) for p in parts], _P())
        try:
            dest = unique_path(out_dir, 'Merged_documents.pdf')
            doc.save(str(dest), garbage=4, deflate=True)
        finally:
            doc.close()
    _emit(progress, 100)
    return [dest]


def docs_pdf_each(paths, out_dir, progress=None, ui_lang='en'):
    """Each Word / PowerPoint / Excel / image file → its own PDF next to the others."""
    global _LANG
    _LANG = ui_lang
    import pymupdf as fitz
    from office_engine import office_to_pdf, OFFICE_EXTS
    from pdf_merge import merge_documents
    outputs = []
    for i, p in enumerate(paths, 1):
        src = Path(p)
        ext = src.suffix.lower()
        if ext in OFFICE_EXTS or ext in SHEET_EXTS:
            outputs += office_to_pdf([src], out_dir)
        elif ext in IMAGE_EXTS:
            class _P:
                def emit(self, v):
                    pass
            doc = merge_documents([str(src)], _P())
            try:
                dest = unique_path(out_dir, src.stem + '.pdf')
                doc.save(str(dest), garbage=4, deflate=True)
            finally:
                doc.close()
            outputs.append(dest)
        # PDFs are already PDFs: nothing to convert.
        _emit(progress, 100 * i / len(paths))
    if not outputs:
        raise ValueError(_t('The selected files are already PDFs.'))
    return outputs


def excel_pdf_fit(paths, out_dir, progress=None, ui_lang='en'):
    """Excel → PDF with every sheet scaled onto a single page (no cut-off columns)."""
    global _LANG
    _LANG = ui_lang
    from office_engine import office_to_pdf
    out = office_to_pdf(paths, out_dir, progress,
                        export='pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}')
    return out


# --------------------------------------------------------------------------- Excel
def _as_xlsx(src, workdir):
    """openpyxl reads .xlsx/.xlsm; older .xls / .ods are converted first."""
    ext = src.suffix.lower()
    if ext in ('.xlsx', '.xlsm'):
        return src
    from office_engine import office_convert
    return office_convert(src, workdir, 'xlsx')


def excel_to_csv(paths, out_dir, progress=None, ui_lang='en'):
    """One UTF-8 CSV per sheet (opens correctly in Excel, Arabic included)."""
    global _LANG
    _LANG = ui_lang
    import openpyxl
    outputs = []
    with tempfile.TemporaryDirectory(prefix='morphix_xl_') as tmp:
        for i, p in enumerate(paths, 1):
            src = Path(p)
            wb = openpyxl.load_workbook(_as_xlsx(src, tmp), data_only=True, read_only=True)
            try:
                single = len(wb.worksheets) == 1
                for ws in wb.worksheets:
                    name = src.stem if single else f'{src.stem} - {_safe(ws.title)}'
                    dest = unique_path(out_dir, name + '.csv')
                    with open(dest, 'w', newline='', encoding='utf-8-sig') as fh:
                        writer = csv.writer(fh)
                        for row in ws.iter_rows(values_only=True):
                            writer.writerow(['' if v is None else v for v in row])
                    outputs.append(dest)
            finally:
                wb.close()
            _emit(progress, 100 * i / len(paths))
    return outputs


def excel_split_sheets(paths, out_dir, progress=None, ui_lang='en'):
    """Each sheet → its own workbook, keeping formatting, widths and formulas."""
    global _LANG
    _LANG = ui_lang
    import openpyxl
    outputs = []
    with tempfile.TemporaryDirectory(prefix='morphix_xl_') as tmp:
        for i, p in enumerate(paths, 1):
            src = _as_xlsx(Path(p), tmp)
            names = openpyxl.load_workbook(src, read_only=True).sheetnames
            if len(names) < 2:
                raise ValueError(_t('{name} has only one sheet.', name=Path(p).name))
            for title in names:
                wb = openpyxl.load_workbook(src)
                for other in list(wb.sheetnames):
                    if other != title:
                        del wb[other]
                dest = unique_path(out_dir, f'{Path(p).stem} - {_safe(title)}.xlsx')
                wb.save(dest)
                outputs.append(dest)
            _emit(progress, 100 * i / len(paths))
    return outputs


def excel_merge(paths, out_dir, progress=None, ui_lang='en'):
    """Several workbooks → one workbook (each sheet copied with its formatting)."""
    global _LANG
    _LANG = ui_lang
    import openpyxl
    from copy import copy
    target = openpyxl.Workbook()
    target.remove(target.active)
    used = set()
    with tempfile.TemporaryDirectory(prefix='morphix_xl_') as tmp:
        for i, p in enumerate(paths, 1):
            src = Path(p)
            if src.suffix.lower() == '.csv':
                ws_new = target.create_sheet(_unique_title(src.stem, used))
                with open(src, newline='', encoding='utf-8-sig', errors='replace') as fh:
                    for row in csv.reader(fh):
                        ws_new.append([_number(v) for v in row])
                continue
            wb = openpyxl.load_workbook(_as_xlsx(src, tmp))
            many = len(wb.worksheets) > 1
            for ws in wb.worksheets:
                title = f'{src.stem[:14]} - {ws.title}' if many else src.stem
                new = target.create_sheet(_unique_title(title, used))
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.__class__.__name__ == 'MergedCell':
                            continue
                        nc = new.cell(row=cell.row, column=cell.column, value=cell.value)
                        if cell.has_style:
                            nc.font, nc.fill, nc.border = copy(cell.font), copy(cell.fill), copy(cell.border)
                            nc.alignment, nc.protection = copy(cell.alignment), copy(cell.protection)
                            nc.number_format = cell.number_format
                for key, dim in ws.column_dimensions.items():
                    new.column_dimensions[key].width = dim.width
                for key, dim in ws.row_dimensions.items():
                    new.row_dimensions[key].height = dim.height
                for rng in ws.merged_cells.ranges:
                    new.merge_cells(str(rng))
                new.sheet_view.rightToLeft = ws.sheet_view.rightToLeft
                new.freeze_panes = ws.freeze_panes
            wb.close()
            _emit(progress, 100 * i / len(paths))
    dest = unique_path(out_dir, 'Merged_workbook.xlsx')
    target.save(dest)
    return [dest]


def csv_to_excel(paths, out_dir, progress=None, ui_lang='en'):
    global _LANG
    _LANG = ui_lang
    import openpyxl
    from openpyxl.utils import get_column_letter
    outputs = []
    for i, p in enumerate(paths, 1):
        src = Path(p)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _safe(src.stem)[:31] or 'Sheet1'
        widths = {}
        with open(src, newline='', encoding='utf-8-sig', errors='replace') as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            for row in csv.reader(fh, dialect):
                ws.append([_number(v) for v in row])
                for c, v in enumerate(row, 1):
                    widths[c] = max(widths.get(c, 0), min(60, len(str(v)) + 2))
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = max(8, w)
        ws.freeze_panes = 'A2'
        if any('؀' <= ch <= 'ۿ' for ch in sample):
            ws.sheet_view.rightToLeft = True      # Arabic data reads right-to-left
        dest = unique_path(out_dir, src.stem + '.xlsx')
        wb.save(dest)
        outputs.append(dest)
        _emit(progress, 100 * i / len(paths))
    return outputs


def _number(value):
    v = value.strip()
    if not v:
        return None
    try:
        if v.isdigit() and len(v) > 1 and v[0] == '0':
            return v                       # keep leading zeros (phone numbers, IDs)
        return int(v) if v.lstrip('-').isdigit() else float(v)
    except ValueError:
        return value


def _safe(name):
    return ''.join('_' if c in '\\/:*?"<>|[]' else c for c in str(name)).strip() or 'Sheet'


def _unique_title(title, used):
    base = _safe(title)[:31] or 'Sheet'
    candidate, n = base, 2
    while candidate.lower() in used:
        candidate = f'{base[:28]} {n}'
        n += 1
    used.add(candidate.lower())
    return candidate
