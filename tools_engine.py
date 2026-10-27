"""Morphix 1.9 tools (headless, run in the worker process).

* PowerPoint-compatible video / audio (H.264 + AAC MP4, AAC M4A)
* Trim video / audio
* HEIC (iPhone photos) → JPG
* PDF → Word (.docx)
* Organize PDF pages (order, rotation, removal)
* Sign PDFs and images
"""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from core_ops import unique_path

VIDEO_PPT_HEIGHTS = {'keep': None, '1080': 1080, '720': 720}


# --------------------------------------------------------------------- FFmpeg
def _ffmpeg():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


_DURATION = re.compile(r'Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)')


def media_info(path):
    """Duration in seconds and whether the file has video / audio streams."""
    proc = subprocess.run([_ffmpeg(), '-hide_banner', '-i', str(path)], capture_output=True,
                          text=True, errors='replace')
    err = proc.stderr
    m = _DURATION.search(err)
    duration = (int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3))) if m else 0.0
    has_video = bool(re.search(r'Stream #\S+.*Video:(?!.*(?:mjpeg|png).*attached pic)', err))
    if re.search(r'Video: (?:mjpeg|png)', err) and 'attached pic' in err:
        has_video = False  # album art in an MP3 / M4A is not a video
    return {'duration': duration, 'video': has_video, 'audio': 'Audio:' in err}


def run_ffmpeg(args, duration=0.0, progress=None, base=0, span=100):
    """Run FFmpeg, reporting progress from its -progress output."""
    cmd = [_ffmpeg(), '-hide_banner', '-loglevel', 'error', '-nostdin', '-y', '-progress', 'pipe:1', *args]
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, errors='replace')
    last = -1
    for line in proc.stdout:
        if progress and duration > 0 and line.startswith('out_time_us='):
            try:
                done = int(line.split('=', 1)[1]) / 1_000_000
            except ValueError:
                continue
            value = base + int(min(1.0, done / duration) * span)
            if value != last:
                last = value
                progress(min(99, value))
    err = proc.stderr.read()
    if proc.wait():
        raise RuntimeError(err.strip()[-1500:] or 'FFmpeg failed')


def ppt_media(paths, out_dir, progress=None, *, max_height='1080', audio_format='m4a'):
    """Re-encode to what PowerPoint plays everywhere: MP4 (H.264 + AAC) / M4A (AAC)."""
    outputs = []
    for i, src in enumerate(paths):
        src = Path(src)
        info = media_info(src)
        base, span = i * 100 // len(paths), 100 // len(paths)
        if info['video']:
            out = unique_path(out_dir, src.stem + '_PowerPoint.mp4')
            height = VIDEO_PPT_HEIGHTS.get(str(max_height))
            scale = (f"scale=-2:'min({height},ih)':flags=lanczos" if height
                     else "scale='trunc(iw/2)*2':'trunc(ih/2)*2'")
            args = ['-i', str(src), '-map', '0:v:0', '-map', '0:a:0?',
                    '-vf', scale,
                    '-c:v', 'libx264', '-profile:v', 'high', '-level:v', '4.2', '-pix_fmt', 'yuv420p',
                    '-preset', 'medium', '-crf', '20',
                    '-c:a', 'aac', '-b:a', '192k', '-ac', '2', '-ar', '48000',
                    '-movflags', '+faststart', str(out)]
            run_ffmpeg(args, info['duration'], progress, base, span)
        elif info['audio']:
            if audio_format == 'mp3':
                out = unique_path(out_dir, src.stem + '_PowerPoint.mp3')
                codec = ['-c:a', 'libmp3lame', '-b:a', '192k']
            else:
                out = unique_path(out_dir, src.stem + '_PowerPoint.m4a')
                codec = ['-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart']
            run_ffmpeg(['-i', str(src), '-vn', '-ac', '2', '-ar', '48000', *codec, str(out)],
                       info['duration'], progress, base, span)
        else:
            raise ValueError(f'No video or audio found in {src.name}')
        outputs.append(out)
    if progress:
        progress(100)
    return outputs


def parse_time(value):
    """'90', '1:30', '01:02:03.5' → seconds."""
    text = str(value).strip().replace(',', '.')
    if not text:
        return 0.0
    parts = text.split(':')
    if len(parts) > 3:
        raise ValueError(f'Invalid time: {value}')
    seconds = 0.0
    for part in parts:
        seconds = seconds * 60 + float(part)
    if seconds < 0:
        raise ValueError(f'Invalid time: {value}')
    return seconds


def format_time(seconds):
    seconds = max(0.0, float(seconds))
    h, rest = divmod(seconds, 3600)
    m, s = divmod(rest, 60)
    return f'{int(h):02d}:{int(m):02d}:{s:06.3f}'


AUDIO_ENCODERS = {'.mp3': ['-c:a', 'libmp3lame', '-q:a', '2'], '.m4a': ['-c:a', 'aac', '-b:a', '192k'],
                  '.aac': ['-c:a', 'aac', '-b:a', '192k'], '.wav': ['-c:a', 'pcm_s16le'],
                  '.flac': ['-c:a', 'flac'], '.ogg': ['-c:a', 'libvorbis', '-q:a', '6'],
                  '.opus': ['-c:a', 'libopus', '-b:a', '128k']}


def trim_media(paths, out_dir, progress=None, *, start='0', end='', exact=True):
    """Keep only start → end. 'exact' re-encodes (frame-accurate); otherwise a fast stream copy."""
    outputs = []
    t0 = parse_time(start)
    for i, src in enumerate(paths):
        src = Path(src)
        info = media_info(src)
        t1 = parse_time(end) if str(end).strip() else info['duration']
        if info['duration'] and t1 > info['duration']:
            t1 = info['duration']
        if t1 <= t0:
            raise ValueError('The end time must be after the start time.')
        length = t1 - t0
        ext = src.suffix.lower()
        stamp = '_{}-{}s'.format(*(f'{v:.1f}'.rstrip('0').rstrip('.') for v in (t0, t1)))
        if info['video']:
            if exact:
                out = unique_path(out_dir, src.stem + stamp + '.mp4')
                args = ['-ss', f'{t0:.3f}', '-i', str(src), '-t', f'{length:.3f}', '-map', '0:v:0', '-map', '0:a:0?',
                        '-c:v', 'libx264', '-preset', 'medium', '-crf', '18', '-pix_fmt', 'yuv420p',
                        '-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart', str(out)]
            else:
                keep = ext if ext in ('.mp4', '.mov', '.m4v', '.mkv', '.webm') else '.mkv'
                out = unique_path(out_dir, src.stem + stamp + keep)
                args = ['-ss', f'{t0:.3f}', '-i', str(src), '-t', f'{length:.3f}', '-map', '0',
                        '-c', 'copy', '-avoid_negative_ts', 'make_zero', str(out)]
        elif info['audio']:
            keep = ext if ext in AUDIO_ENCODERS else '.m4a'
            out = unique_path(out_dir, src.stem + stamp + keep)
            codec = AUDIO_ENCODERS[keep] if exact else ['-c:a', 'copy']
            if not exact and ext not in AUDIO_ENCODERS:
                codec = AUDIO_ENCODERS[keep]
            args = ['-ss', f'{t0:.3f}', '-i', str(src), '-t', f'{length:.3f}', '-vn', *codec, str(out)]
        else:
            raise ValueError(f'No video or audio found in {src.name}')
        run_ffmpeg(args, length, progress, i * 100 // len(paths), 100 // len(paths))
        outputs.append(out)
    if progress:
        progress(100)
    return outputs


# --------------------------------------------------------------------- HEIC
def register_heic():
    """Teach Pillow to open HEIC / HEIF / AVIF (iPhone photos). Safe to call often."""
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
        return True
    except Exception:
        return False


def heic_to_jpg(paths, out_dir, progress=None, *, quality=92):
    from PIL import Image, ImageOps
    if not register_heic():
        raise RuntimeError('HEIC support is not installed in this build.')
    outputs = []
    for i, src in enumerate(paths, 1):
        src = Path(src)
        with Image.open(src) as original:
            exif = original.info.get('exif')
            icc = original.info.get('icc_profile')
            image = ImageOps.exif_transpose(original).convert('RGB')
        out = unique_path(out_dir, src.stem + '.jpg')
        params = {'quality': int(quality), 'optimize': True, 'subsampling': 0 if quality >= 90 else 2}
        if icc:
            params['icc_profile'] = icc
        if exif:
            try:
                from PIL import Image as _I
                data = _I.Exif()
                data.load(exif)
                data[0x0112] = 1  # already rotated upright
                params['exif'] = data.tobytes()
            except Exception:
                pass
        image.save(out, 'JPEG', **params)
        outputs.append(out)
        if progress:
            progress(round(i * 100 / len(paths)))
    return outputs


# --------------------------------------------------------------------- PDF → Word
def pdf_to_word(paths, out_dir, progress=None, *, status=None):
    import logging
    import warnings as _w
    import pymupdf as fitz
    _w.filterwarnings('ignore')
    from pdf2docx import Converter
    logging.getLogger('pdf2docx').setLevel(logging.ERROR)
    logging.disable(logging.INFO)
    outputs, warnings = [], []
    for i, src in enumerate(paths):
        src = Path(src)
        with fitz.open(src) as doc:
            if doc.needs_pass:
                raise RuntimeError(f'Password-protected PDF: {src.name}')
            pages = len(doc)
            text_pages = sum(1 for page in doc if page.get_text('text').strip())
        if text_pages == 0:
            warnings.append(src.name)
        out = unique_path(out_dir, src.stem + '.docx')
        conv = Converter(str(src))
        try:
            conv.convert(str(out), start=0, end=None, multi_processing=False)
        finally:
            conv.close()
        try:
            _mark_rtl_paragraphs(out)
        except Exception:
            pass
        if progress:
            progress(round((i + 1) * 100 / len(paths)))
        outputs.append(out)
    if warnings and status:
        status('Scanned PDF without text: run OCR first for editable text — ' + ', '.join(warnings))
    return outputs


_ARABIC = re.compile(r'[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]')
_LATIN = re.compile(r'[A-Za-z]')


def _mark_rtl_paragraphs(docx_path):
    """Arabic paragraphs read right-to-left in Word (bidi + right alignment)."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(str(docx_path))
    def paragraphs():
        yield from doc.paragraphs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
    changed = False
    for par in paragraphs():
        text = par.text
        if len(_ARABIC.findall(text)) <= len(_LATIN.findall(text)):
            continue
        p_pr = par._p.get_or_add_pPr()
        if p_pr.find(qn('w:bidi')) is None:
            bidi = OxmlElement('w:bidi')
            bidi.set(qn('w:val'), '1')
            p_pr.insert(0, bidi)
        for run in par.runs:
            r_pr = run._r.get_or_add_rPr()
            if r_pr.find(qn('w:rtl')) is None:
                rtl = OxmlElement('w:rtl')
                rtl.set(qn('w:val'), '1')
                r_pr.append(rtl)
        changed = True
    if changed:
        doc.save(str(docx_path))


# --------------------------------------------------------------------- organize pages
def organize_pdf(path, out_dir, pages, progress=None):
    """pages: [[original_index, extra_rotation], …] in the new order."""
    import pymupdf as fitz
    src = Path(path)
    if not pages:
        raise ValueError('A PDF must contain at least one page.')
    with fitz.open(src) as doc:
        result = fitz.open()
        try:
            for n, (index, rotation) in enumerate(pages, 1):
                result.insert_pdf(doc, from_page=int(index), to_page=int(index))
                page = result[-1]
                page.set_rotation((page.rotation + int(rotation)) % 360)
                if progress and n % 10 == 0:
                    progress(round(n * 90 / len(pages)))
            out = unique_path(out_dir, src.stem + '_organized.pdf')
            result.save(str(out), garbage=4, deflate=True)
        finally:
            result.close()
    if progress:
        progress(100)
    return [out]


# --------------------------------------------------------------------- signatures
def _page_list(spec, total, current=0):
    from engine_core import parse_pages
    spec = str(spec or 'current').strip().lower()
    if spec == 'all':
        return list(range(total))
    if spec == 'first':
        return [0]
    if spec == 'last':
        return [total - 1]
    if spec == 'current':
        return [max(0, min(total - 1, int(current)))]
    return parse_pages(spec, total)


def sign_pdf(paths, out_dir, progress=None, *, signature, box, pages='current', page_index=0):
    """Place the signature image on PDF pages.

    box = [x, y, w] as fractions of the displayed page (top-left origin);
    the height follows the signature's aspect ratio.
    """
    import pymupdf as fitz
    from PIL import Image
    with Image.open(signature) as im:
        ratio = im.height / max(1, im.width)
    data = Path(signature).read_bytes()
    outputs = []
    for i, src in enumerate(paths):
        src = Path(src)
        with fitz.open(src) as doc:
            if doc.needs_pass:
                raise RuntimeError(f'Password-protected PDF: {src.name}')
            for index in _page_list(pages, len(doc), page_index):
                page = doc[index]
                shown = page.rect  # what the reader sees (rotation applied)
                x, y, w = (float(v) for v in box[:3])
                width = w * shown.width
                rect = fitz.Rect(shown.x0 + x * shown.width, shown.y0 + y * shown.height,
                                 shown.x0 + x * shown.width + width,
                                 shown.y0 + y * shown.height + width * ratio)
                target = rect * page.derotation_matrix
                target.normalize()
                page.insert_image(target, stream=data, keep_proportion=True, overlay=True,
                                  rotate=page.rotation)
            out = unique_path(out_dir, src.stem + '_signed.pdf')
            doc.save(str(out), garbage=3, deflate=True)
        outputs.append(out)
        if progress:
            progress(round((i + 1) * 100 / len(paths)))
    return outputs


def sign_images(paths, out_dir, progress=None, *, signature, box):
    """Same relative position and size on every selected image."""
    from PIL import Image, ImageOps
    register_heic()
    with Image.open(signature) as s:
        sig = s.convert('RGBA')
    outputs = []
    for i, src in enumerate(paths, 1):
        src = Path(src)
        with Image.open(src) as original:
            icc = original.info.get('icc_profile')
            image = ImageOps.exif_transpose(original)
            image = image.convert('RGBA')
        x, y, w = (float(v) for v in box[:3])
        width = max(1, int(round(w * image.width)))
        height = max(1, int(round(width * sig.height / sig.width)))
        stamp = sig.resize((width, height), Image.Resampling.LANCZOS)
        image.alpha_composite(stamp, (int(round(x * image.width)), int(round(y * image.height))))
        ext = src.suffix.lower()
        if ext in ('.jpg', '.jpeg', '.heic', '.heif'):
            out = unique_path(out_dir, src.stem + '_signed.jpg')
            params = {'quality': 95, 'subsampling': 0}
            if icc:
                params['icc_profile'] = icc
            image.convert('RGB').save(out, 'JPEG', **params)
        else:
            keep = ext if ext in ('.png', '.webp', '.tif', '.tiff') else '.png'
            out = unique_path(out_dir, src.stem + '_signed' + keep)
            params = {'icc_profile': icc} if icc else {}
            if keep == '.webp':
                params['lossless'] = True
            image.save(out, **params)
        outputs.append(out)
        if progress:
            progress(round(i * 100 / len(paths)))
    return outputs


def clean_signature(src, dest, *, threshold=200, color=None):
    """Turn a photo / scan of a signature on paper into a transparent PNG.

    Pixels lighter than `threshold` become transparent with a soft edge;
    `color` (hex) optionally recolours the ink.
    """
    from PIL import Image, ImageOps
    import numpy as np
    register_heic()
    with Image.open(src) as original:
        image = ImageOps.exif_transpose(original).convert('RGBA')
    arr = np.asarray(image).astype(np.float32)
    lum = arr[..., :3] @ np.array([0.299, 0.587, 0.114], dtype=np.float32)
    soft = 40.0
    alpha = np.clip((threshold - lum) / soft + 0.5, 0, 1) * (arr[..., 3] / 255.0)
    out = arr.copy()
    if color:
        c = color.lstrip('#')
        out[..., 0], out[..., 1], out[..., 2] = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    out[..., 3] = alpha * 255
    result = Image.fromarray(out.astype(np.uint8), 'RGBA')
    bbox = result.getchannel('A').point(lambda v: 255 if v > 12 else 0).getbbox()
    if bbox:
        pad = 8
        bbox = (max(0, bbox[0] - pad), max(0, bbox[1] - pad), min(result.width, bbox[2] + pad),
                min(result.height, bbox[3] + pad))
        result = result.crop(bbox)
    result.save(dest)
    return Path(dest)


def trim_transparent(src, dest=None, pad=8):
    from PIL import Image
    with Image.open(src) as im:
        im = im.convert('RGBA')
        bbox = im.getchannel('A').point(lambda v: 255 if v > 12 else 0).getbbox()
        if bbox:
            bbox = (max(0, bbox[0] - pad), max(0, bbox[1] - pad), min(im.width, bbox[2] + pad),
                    min(im.height, bbox[3] + pad))
            im = im.crop(bbox)
        im.save(dest or src)
    return Path(dest or src)
